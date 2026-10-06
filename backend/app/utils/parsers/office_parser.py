"""
Office 文档解析：.docx / .pptx / .xlsx / .xlsm
--------------------------------------------
共同目标：保留文档结构，而不是拍平成一大段文字 —— 本体生成要靠标题层级、
工作表名、页边界来判断"这段话在说什么"。

各解析库一律在函数内惰性 import：缺失时降级成逐文件错误串，而不是让整个
Flask 应用 import 失败或返回 500。

旧版二进制格式（.doc / .ppt / .xls）不在这里处理 —— 它们由 file_parser 的
classify() 拦下并给出「另存为 .docx/.pptx/.xlsx」的提示。
"""

import logging
import re
from typing import Any, Iterator, List, Optional, Tuple

from ..locale import t
from . import base
from .base import ExtractionResult

logger = logging.getLogger(__name__)

_DOCX_BLOCK_TAGS = ('}p', '}tbl')
_DOCX_TEXT_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'


# ============================ .docx ============================

def extract_docx(path: str) -> ExtractionResult:
    """.docx —— 按文档顺序输出标题/段落/表格，末尾汇总文本框与页眉页脚。"""
    base.ensure_readable_package(path, 'zip')

    try:
        from docx import Document
    except ImportError as exc:
        raise base.ExtractionError('需要安装 python-docx: uv sync') from exc

    doc = Document(path)

    blocks: List[str] = []
    notes: List[str] = []
    block_count = 0

    for kind, obj in _iter_docx_blocks(doc):
        if block_count >= base.MAX_DOCX_BLOCKS:
            notes.append(f'block-cap:{base.MAX_DOCX_BLOCKS}')
            blocks.append(base.note(t('parse.blocksTruncated', shown=base.MAX_DOCX_BLOCKS)))
            break
        block_count += 1
        if kind == 'para':
            md = _docx_para_to_md(obj)
        else:
            md = _docx_table_to_md(obj)
        if md:
            blocks.append(md)

    textboxes = _docx_textboxes(doc)
    if textboxes:
        blocks.append(base.join_blocks([
            base.heading(2, t('parse.textbox')),
            '\n'.join(f'- {tb}' for tb in textboxes),
        ]))

    headers, footers = _docx_headers_footers(doc)
    for label_key, items in (('parse.header', headers), ('parse.footer', footers)):
        if items:
            blocks.append(base.join_blocks([
                base.heading(2, t(label_key)),
                '\n'.join(f'- {item}' for item in items),
            ]))

    return ExtractionResult(text=base.join_blocks(blocks), notes=notes)


def _iter_docx_blocks(doc) -> Iterator[Tuple[str, Any]]:
    """
    按文档顺序产出 ('para'|'table', obj)。

    必须走 body.iterchildren()：doc.paragraphs 与 doc.tables 是两个互相独立、
    丢失交错的列表，直接拼接会把所有表格挪到文末，破坏段落上下文。
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        tag = child.tag
        if not isinstance(tag, str):
            continue
        if tag.endswith(_DOCX_BLOCK_TAGS[0]):
            yield 'para', Paragraph(child, doc)
        elif tag.endswith(_DOCX_BLOCK_TAGS[1]):
            yield 'table', Table(child, doc)


def _docx_para_to_md(paragraph) -> str:
    """按段落样式还原标题层级；列表统一压成单层（不再缩进）。"""
    text = (paragraph.text or '').strip()
    if not text:
        return ''

    try:
        style_name = paragraph.style.name if paragraph.style is not None else ''
    except Exception:
        style_name = ''
    style_name = style_name or ''

    if style_name == 'Title':
        return base.heading(1, text)
    if style_name == 'Subtitle':
        return base.heading(2, text)

    match = re.match(r'^Heading (\d+)$', style_name)
    if match:
        return base.heading(int(match.group(1)), text)

    if style_name.startswith('List Bullet') or style_name.startswith('List Number'):
        return f'- {text}'

    return text


def _docx_table_to_md(table) -> str:
    """docx 表格 → Markdown。合并单元格在 python-docx 里会重复同一格文本，如实保留。"""
    try:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
    except Exception as exc:
        logger.warning('读取 docx 表格失败: %s', exc)
        return ''
    return base.render_markdown_table(
        rows,
        max_rows=base.MAX_TABLE_ROWS,
        max_cols=base.TABLE_COL_CAP,
    )


def _docx_textboxes(doc) -> List[str]:
    """
    文本框内容。python-docx 不建模文本框，只能从 XML 里捞 w:txbxContent。

    位置信息会丢失，统一汇总到文末 —— 对实体抽取来说，文本本身比位置更重要。
    """
    found: List[str] = []
    try:
        for node in doc.element.body.iter():
            if not isinstance(node.tag, str) or not node.tag.endswith('}txbxContent'):
                continue
            parts = [t.text for t in node.iter(_DOCX_TEXT_NS) if t.text]
            text = ' '.join(p.strip() for p in parts if p and p.strip())
            if text:
                found.append(text)
    except Exception as exc:
        logger.warning('读取 docx 文本框失败: %s', exc)
    return found


def _docx_headers_footers(doc) -> Tuple[List[str], List[str]]:
    """
    页眉/页脚，跨 section 按文本去重。

    值得单独提取：中文报告的表头常年写着发布机构，那是很有价值的实体来源。
    """
    headers: List[str] = []
    footers: List[str] = []
    seen = set()

    try:
        sections = list(doc.sections)
    except Exception:
        return headers, footers

    for section in sections:
        for target, bucket in ((getattr(section, 'header', None), headers),
                               (getattr(section, 'footer', None), footers)):
            if target is None:
                continue
            try:
                paragraphs = list(target.paragraphs)
            except Exception:
                continue
            for paragraph in paragraphs:
                text = (paragraph.text or '').strip()
                if text and text not in seen:
                    seen.add(text)
                    bucket.append(text)
    return headers, footers


# ============================ .pptx ============================

def extract_pptx(path: str) -> ExtractionResult:
    """.pptx —— 每页一个章节：标题、正文要点、表格/图表、备注。"""
    base.ensure_readable_package(path, 'zip')

    try:
        from pptx import Presentation
    except ImportError as exc:
        raise base.ExtractionError('需要安装 python-pptx: uv sync') from exc

    prs = Presentation(path)
    slides = list(prs.slides)
    total = len(slides)
    shown = slides[:base.MAX_SLIDES]

    blocks: List[str] = []
    for index, slide in enumerate(shown, start=1):
        slide_md = _pptx_slide_to_md(slide, index)
        if slide_md:
            blocks.append(slide_md)

    notes: List[str] = []
    if total > base.MAX_SLIDES:
        notes.append(f'slide-cap:{base.MAX_SLIDES}')
        blocks.append(base.note(t('parse.slidesTruncated', total=total, shown=len(shown))))

    return ExtractionResult(text=base.join_blocks(blocks), notes=notes)


def _iter_shapes(shapes) -> Iterator[Any]:
    """
    递归展开形状。

    顶层 slide.shapes 取不到分组（Group）内部的元素 —— 这是 pptx 解析最常见的漏内容原因。
    """
    try:
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        MSO_SHAPE_TYPE = None

    for shape in shapes:
        yield shape
        is_group = False
        if MSO_SHAPE_TYPE is not None:
            try:
                is_group = shape.shape_type == MSO_SHAPE_TYPE.GROUP
            except Exception:
                is_group = False
        if is_group:
            try:
                yield from _iter_shapes(shape.shapes)
            except Exception:
                continue


def _pptx_title(slide) -> Tuple[str, Optional[Any]]:
    """
    返回 (标题文本, 标题来源形状)。

    来源形状会被调用方跳过，否则标题文字会在正文要点里再出现一遍。
    """
    title_shape = None
    try:
        title_shape = slide.shapes.title
    except Exception:
        title_shape = None

    if title_shape is not None:
        text = (getattr(title_shape, 'text', '') or '').strip()
        if text:
            return text, title_shape

    for shape in _iter_shapes(slide.shapes):
        if getattr(shape, 'has_text_frame', False):
            text = (shape.text_frame.text or '').strip()
            if text:
                return text[:40], shape

    return t('parse.noTitle'), None


def _pptx_slide_to_md(slide, index: int) -> str:
    title, title_source = _pptx_title(slide)
    # 比较底层 XML 元素而不是 Python 对象：python-pptx 每次访问 shapes 都会新建代理
    # 对象，`shape is title_shape` 恒为 False，标题会被当成正文再输出一遍。
    title_element = getattr(title_source, '_element', None)
    parts: List[str] = [base.heading(2, f'Slide {index}: {title}')]

    bullets: List[str] = []
    has_picture = False

    for shape in _iter_shapes(slide.shapes):
        if title_element is not None and getattr(shape, '_element', None) is title_element:
            continue
        try:
            if getattr(shape, 'has_table', False):
                table_md = _pptx_table_to_md(shape.table)
                if table_md:
                    parts.append(table_md)
                continue

            if getattr(shape, 'has_chart', False):
                chart_md = _pptx_chart_to_md(shape.chart)
                if chart_md:
                    parts.append(chart_md)
                continue

            if getattr(shape, 'has_text_frame', False):
                for line in (shape.text_frame.text or '').splitlines():
                    line = line.strip()
                    if line:
                        bullets.append(line)
                continue

            try:
                from pptx.enum.shapes import MSO_SHAPE_TYPE
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    has_picture = True
            except Exception:
                pass
        except Exception as exc:
            logger.warning('解析 pptx 形状失败 (slide %s): %s', index, exc)
            continue

    if bullets:
        parts.append('\n'.join(f'- {line}' for line in bullets))

    # 备注：必须先判 has_notes_slide —— 访问 notes_slide 属性会就地创建备注页，
    # 不加判断会让每份解析过的 deck 都凭空长出空备注
    try:
        if slide.has_notes_slide:
            notes_text = (slide.notes_slide.notes_text_frame.text or '').strip()
            if notes_text:
                note_lines = [line.strip() for line in notes_text.splitlines() if line.strip()]
                parts.append(base.join_blocks([
                    base.heading(3, t('parse.notes')),
                    '\n'.join(f'- {line}' for line in note_lines),
                ]))
    except Exception as exc:
        logger.warning('读取 pptx 备注失败 (slide %s): %s', index, exc)

    if len(parts) == 1:
        # 除了标题什么都没有：要么真空白页，要么是纯图片页
        parts.append(base.note(t('parse.imageOnlySlide') if has_picture else t('parse.emptySlide')))

    return base.join_blocks(parts)


def _pptx_table_to_md(table) -> str:
    try:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
    except Exception as exc:
        logger.warning('读取 pptx 表格失败: %s', exc)
        return ''
    return base.render_markdown_table(
        rows,
        max_rows=base.MAX_TABLE_ROWS,
        max_cols=base.TABLE_COL_CAP,
    )


def _pptx_chart_to_md(chart) -> str:
    """
    图表 → Markdown 表格（类别 + 各系列取值）。

    整段包 try/except：图表 XML 是 OOXML 里最不统一的部分，
    不能因为某个不支持的图表类型就让整份 deck 解析失败。
    """
    try:
        plots = list(chart.plots)
    except Exception:
        return ''
    if not plots:
        return ''

    plot = plots[0]
    try:
        categories = [base.cell_to_str(c) for c in (plot.categories or [])]
        series = list(plot.series)
    except Exception:
        return ''
    if not categories or not series:
        return ''

    header = [t('parse.chartCategory')] + [
        (getattr(s, 'name', None) or t('parse.chartSeries', n=i + 1)) for i, s in enumerate(series)
    ]

    rows: List[List[str]] = []
    series_values = []
    for s in series:
        try:
            series_values.append(list(s.values))
        except Exception:
            series_values.append([])

    for idx, category in enumerate(categories):
        row = [category]
        for values in series_values:
            row.append(base.cell_to_str(values[idx]) if idx < len(values) else '')
        rows.append(row)

    table_md = base.render_markdown_table(
        rows,
        header=header,
        max_rows=base.MAX_TABLE_ROWS,
        max_cols=base.TABLE_COL_CAP,
    )
    if not table_md:
        return ''

    try:
        chart_type = str(chart.chart_type).split()[0]
    except Exception:
        chart_type = ''
    label = t('parse.chartHeading', type=chart_type) if chart_type else t('parse.chartHeadingPlain')
    return base.join_blocks([base.heading(3, label), table_md])


# ============================ .xlsx / .xlsm ============================

def extract_xlsx(path: str) -> ExtractionResult:
    """
    .xlsx / .xlsm —— 每个工作表一个章节。

    data_only=True 取作者当时看到的计算值。若整本工作簿在 data_only 下全空，
    说明公式没有缓存值（脚本生成的文件很常见），改用 data_only=False 重读一次，
    输出公式原文并明确标注。

    已知限制：**合并单元格的取值只出现在左上角，其余格渲染为空**。
    read_only 模式下的 ReadOnlyWorksheet 不暴露 merged_cells，拿不到合并区间，
    也就无法加说明。对本体抽取的影响有限（表头文字仍然在）。
    """
    base.ensure_readable_package(path, 'zip')

    try:
        import openpyxl
    except ImportError as exc:
        raise base.ExtractionError('需要安装 openpyxl: uv sync') from exc

    blocks, notes, has_value = _render_workbook(openpyxl, path, data_only=True)

    if not has_value:
        blocks2, notes2, has_value2 = _render_workbook(openpyxl, path, data_only=False)
        if has_value2:
            blocks2.insert(0, base.note(t('parse.formulaNotCached')))
            return ExtractionResult(text=base.join_blocks(blocks2), notes=notes + notes2)

    return ExtractionResult(text=base.join_blocks(blocks), notes=notes)


def _render_workbook(openpyxl, path: str, data_only: bool) -> Tuple[List[str], List[str], bool]:
    """返回 (块列表, 说明列表, 是否出现过非空值)。"""
    # keep_links=False：阻止 openpyxl 跟随外部工作簿链接
    workbook = openpyxl.load_workbook(
        path, read_only=True, data_only=data_only, keep_links=False,
    )

    blocks: List[str] = []
    notes: List[str] = []
    has_value = False

    try:
        sheet_names = list(workbook.sheetnames)
        if len(sheet_names) > base.MAX_SHEETS:
            notes.append(f'sheet-cap:{base.MAX_SHEETS}')

        for sheet_name in sheet_names[:base.MAX_SHEETS]:
            try:
                worksheet = workbook[sheet_name]
            except Exception as exc:
                logger.warning('打开工作表失败 %s: %s', sheet_name, exc)
                continue

            rows, rows_truncated, cols_truncated = _read_sheet(worksheet)
            if not rows:
                continue
            has_value = True

            header, data_rows = rows[0], rows[1:]
            # 行/列的截断说明由本函数自己给（真实行数在流式读取下拿不到），
            # 所以这里把 render 的上限放到刚好不触发它自己的说明
            table_md = base.render_markdown_table(
                data_rows,
                header=header,
                max_rows=max(len(data_rows), 1),
                max_cols=base.MAX_SHEET_COLS,
            )
            sheet_blocks = [base.heading(2, t('parse.sheetHeading', name=sheet_name))]

            if table_md:
                sheet_blocks.append(table_md)
            else:
                continue

            truncation = []
            if rows_truncated:
                truncation.append(t('parse.rowsTruncated', shown=len(data_rows)))
            if cols_truncated:
                truncation.append(t('parse.colsTruncated', shown=base.MAX_SHEET_COLS))
            if truncation:
                sheet_blocks.append(base.note('；'.join(truncation)))
                notes.append(f'truncated:{sheet_name}')

            blocks.append(base.join_blocks(sheet_blocks))

        if len(sheet_names) > base.MAX_SHEETS:
            blocks.append(base.note(t('parse.sheetsTruncated',
                                      total=len(sheet_names), shown=base.MAX_SHEETS)))
    finally:
        try:
            workbook.close()
        except Exception:
            pass

    return blocks, notes, has_value


def _read_sheet(worksheet) -> Tuple[List[List[Any]], bool, bool]:
    """
    读一个工作表，返回 (行列表, 是否被截断行, 是否可能被截断列)。

    刻意**不信任** ws.max_row / ws.max_column：read_only 模式下它们来自
    `<dimension ref>` 元素，而很多工具把它写成整表（A1:XFD1048576），
    信任它就会把请求挂死。改为按 max_col 读取并计数到上限即停。
    """
    rows: List[List[Any]] = []
    rows_truncated = False
    cols_truncated = False
    header_found = False

    try:
        row_iter = worksheet.iter_rows(max_col=base.MAX_SHEET_COLS, values_only=True)
    except Exception as exc:
        logger.warning('读取工作表行失败: %s', exc)
        return rows, False, False

    for raw_row in row_iter:
        values = list(raw_row)
        if not header_found:
            # 跳过前置的整行空白，用它来定位真正的表头
            if not any(not _is_blank(v) for v in values):
                continue
            header_found = True

        # 表头不占配额：MAX_SHEET_ROWS 约束的是渲染出来的数据行数，
        # 所以 rows 允许到 MAX_SHEET_ROWS + 1（首行是表头）
        if len(rows) > base.MAX_SHEET_ROWS:
            rows_truncated = True
            break

        # 最后一列有值 → 很可能还有被 max_col 截掉的列
        if not cols_truncated and values and not _is_blank(values[-1]):
            cols_truncated = True

        while values and _is_blank(values[-1]):
            values.pop()
        rows.append(values)

    return rows, rows_truncated, cols_truncated


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    return False
